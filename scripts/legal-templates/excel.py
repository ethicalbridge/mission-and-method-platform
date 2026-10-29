"""Excel versions of the register-style templates.
Every pipe table in src/<slug>.md becomes its own worksheet (named after the nearest heading),
after a "How to use" sheet. Yellow cells = fields to fill in, as in the Word templates.
Usage: python3 scripts/legal-templates/excel.py [slug ...]   (no args = the default register set)
Writes assets/legal-templates/<slug>.xlsx
"""
import os, re, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(ROOT, '..', '..', 'assets', 'legal-templates')
DEFAULT = ['risk-register', 'delegation-of-authority-matrix', 'records-retention-schedule', 'board-skills-matrix']
EDITION = '2026-10'
TEAL, INK, CREAM, YELLOW = '16746E', '102E3B', 'FFFAF3', 'FFF2A8'
thin = Side(style='thin', color='D8DFD9')

def clean(t):
    t = t.replace('<br>', '\n').replace('**', '')
    return re.sub(r'\[\[([^\]]+)\]\]', r'[\1]', t).strip()

def meta(src, key):
    m = re.search(rf'^{key}:\s*(.*)$', src, re.M)
    return clean(m.group(1)) if m else ''

def tables(src):
    body = src.split('\n---\n', 1)[1] if src.startswith('---') else src
    lines, out, heading, i = body.split('\n'), [], 'Table', 0
    while i < len(lines):
        ln = lines[i]
        if re.match(r'^#{1,4} ', ln):
            heading = re.sub(r'^#+\s*', '', ln).strip()
        if ln.startswith('|'):
            block = []
            while i < len(lines) and lines[i].startswith('|'):
                block.append(lines[i]); i += 1
            rows = [[c.strip() for c in r.strip().strip('|').split('|')] for r in block]
            rows = [r for r in rows if not all(re.fullmatch(r':?-{2,}:?', c) for c in r if c)]
            if len(rows) >= 2:
                out.append((heading, rows))
            continue
        i += 1
    return out

def sheet_name(h, used):
    base = re.sub(r'[\[\]\*\?/\\:]', '', clean(h))[:28] or 'Table'
    name, n = base, 2
    while name in used:
        name = f'{base[:25]} {n}'; n += 1
    used.add(name); return name

def build(slug):
    src = open(os.path.join(ROOT, 'src', slug + '.md'), encoding='utf8').read()
    title = meta(src, 'title')
    wb = Workbook(); ws = wb.active; ws.title = 'How to use'
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 110
    rows = [(title, Font(name='Georgia', size=20, color=INK)),
            ('Method into Impact · Excel version · template edition ' + EDITION, Font(bold=True, size=10, color=TEAL)),
            ('', None),
            (meta(src, 'summary'), Font(size=11, color=INK)),
            ('', None),
            ('How to use this file', Font(bold=True, size=12, color=INK)),
            ('• Each tab is one table from the Word template. The Word version has the full guidance and the "Adapt this" notes.', Font(size=11)),
            ('• Yellow cells are fields to fill in. Rows marked as examples show the level of detail expected: replace or delete them.', Font(size=11)),
            ('• Add rows as you need them. Keep this file with your board papers and review it at least once a year.', Font(size=11)),
            ('', None),
            ('This is a template, not legal advice. Adapt it to your organisation and check it against the law where you operate. Licensed for use within one organisation; please do not resell or republish it.', Font(italic=True, size=10, color='5D6B73'))]
    for r, (txt, font) in enumerate(rows, 1):
        c = ws.cell(row=r, column=1, value=txt); c.alignment = Alignment(wrap_text=True, vertical='top')
        if font: c.font = font
    used = {'How to use'}
    for heading, data in tables(src):
        sh = wb.create_sheet(sheet_name(heading, used))
        sh.sheet_view.showGridLines = False
        ncol = max(len(r) for r in data)
        widths = [12] * ncol
        for ri, row in enumerate(data, 1):
            for ci in range(ncol):
                raw = row[ci] if ci < len(row) else ''
                val = clean(raw)
                cell = sh.cell(row=ri, column=ci + 1, value=val)
                cell.alignment = Alignment(wrap_text=True, vertical='top')
                cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
                if ri == 1:
                    cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor=TEAL)
                elif '[[' in raw:
                    cell.fill = PatternFill('solid', fgColor=YELLOW)
                longest = max((len(x) for x in val.split('\n')), default=0)
                widths[ci] = max(widths[ci], min(60, longest + 2))
        for ci, w in enumerate(widths, 1):
            sh.column_dimensions[get_column_letter(ci)].width = w
        sh.freeze_panes = 'A2'
        sh.auto_filter.ref = f'A1:{get_column_letter(ncol)}{len(data)}'
    path = os.path.join(OUT, slug + '.xlsx')
    wb.save(path)
    print(f'built {slug}.xlsx ({len(wb.sheetnames) - 1} tables)')

if __name__ == '__main__':
    for s in (sys.argv[1:] or DEFAULT):
        build(s)
